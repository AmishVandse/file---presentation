import streamlit as st
from io import StringIO
import openai
from fpdf import FPDF
import fitz  # PyMuPDF
import io
import docx


openai.api_key = st.text_input("Enter API key", type="password")

def extract_text_from_pdf(pdf_file):
    doc = fitz.open(stream=pdf_file.read(), filetype="pdf")
    text = ""
    for page_num in range(len(doc)):
        page = doc.load_page(page_num)
        text += page.get_text()
    return text

def extract_text_from_docx(docx_file):
    doc = docx.Document(docx_file)
    text = "\n".join([para.text for para in doc.paragraphs])
    return text

def extract_text_from_txt(txt_file):
    stringio = StringIO(txt_file.getvalue().decode("utf-8"))
    text = stringio.read()
    return text

def summarize_text(doc_names, context):
    prompt = (
        "You are proficient in document analysis, cybersecurity research, and artificial intelligence. Your writing style is precise rather than overly embellished. Upon receiving a collection of documents, each listed with a name in quotes, followed by their content delimited like XML, you will extract pertinent information to craft a cohesive presentation. This presentation will integrate the main ideas from all documents, focusing on the applications of Generative AI in Cybersecurity. Proper citations from the documents will be included throughout the Markdown-based Point-of-View presentation."
        f"{doc_names}\n\n{context}"
    )
    response = openai.chat.completions.create(
        model="gpt-3.5-turbo",
        messages=[
            {"role": "system", "content": "You are a helpful assistant."},
            {"role": "user", "content": prompt}
        ],
        max_tokens=1500
    )
    return response.choices[0].message['content'].strip()

class PDF(FPDF):
    def header(self):
        self.set_font("Arial", "B", 12)
        self.cell(0, 10, "PDF Presentation", 0, 1, "C")

    def footer(self):
        self.set_y(-15)
        self.set_font("Arial", "I", 8)
        self.cell(0, 10, f"Page {self.page_no()}", 0, 0, "C")

    def chapter_title(self, title):
        self.set_font("Arial", "B", 12)
        self.cell(0, 10, title, 0, 1, "L")
        self.ln(10)

    def chapter_body(self, body):
        self.set_font("Arial", "", 12)
        self.multi_cell(0, 10, body)
        self.ln()

def create_presentation_pdf(summaries):
    pdf = PDF()
    pdf.add_page()
    for i, summary in enumerate(summaries):
        pdf.chapter_title(f"Summary {i+1}")
        pdf.chapter_body(summary)

    buffer = io.BytesIO()
    pdf.output(buffer)
    buffer.seek(0)
    return buffer

def main():
    st.title("File to Presentation")

    uploaded_files = st.file_uploader("Upload PDF, DOCX, or TXT files", type=["pdf", "docx", "txt"], accept_multiple_files=True)

    if uploaded_files and openai.api_key:
        context_map = {}
        doc_names = ""
        for uploaded_file in uploaded_files:
            st.subheader(f"Processing: {uploaded_file.name}")
            if uploaded_file.name.endswith(".pdf"):
                text = extract_text_from_pdf(uploaded_file)
            elif uploaded_file.name.endswith(".docx"):
                text = extract_text_from_docx(uploaded_file)
            elif uploaded_file.name.endswith(".txt"):
                text = extract_text_from_txt(uploaded_file)
            else:
                st.error(f"Unsupported file type: {uploaded_file.name}")
                continue
            
            context_map[uploaded_file.name] = text
            doc_names += f'"{uploaded_file.name}"\n'
            st.text("Extracted Text:")
            st.write(text)

        if st.button("Summarize and Generate Presentation"):
            context = "".join([f"<{name}> {content} </{name}>\n" for name, content in context_map.items()])
            summary = summarize_text(doc_names, context)
            st.text("Generated Presentation:")
            st.write(summary)

            presentation_pdf = create_presentation_pdf([summary])
            st.download_button(
                label="Download Presentation PDF",
                data=presentation_pdf,
                file_name="presentation.pdf",
                mime="application/pdf"
            )

if __name__ == "__main__":
    main()
